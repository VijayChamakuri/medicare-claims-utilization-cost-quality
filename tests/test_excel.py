"""Workbook structure, formulas, values, and reconciliation to DuckDB."""

import re
import zipfile

import pytest
from openpyxl import load_workbook

from medicare_claims.excel import HEADER_ROW, SHEETS, build_workbook
from medicare_claims.metrics import compute_kpis


@pytest.fixture(scope="module")
def workbook(warehouse, config, tmp_path_factory):
    path = build_workbook(warehouse, config, tmp_path_factory.mktemp("xl") / "claims.xlsx")
    return path, load_workbook(path)


def test_nine_sheets_in_the_specified_order(workbook) -> None:
    assert workbook[1].sheetnames == SHEETS == ["Executive_Summary", "Monthly_Utilization", "Monthly_Payments", "Provider_Review",
                                                "Readmission_Review", "Member_Risk_Tiers", "Data_Quality", "Metric_Dictionary", "Reconciliation"]


def test_every_sheet_carries_the_synthetic_banner_and_a_table(workbook) -> None:
    for ws in workbook[1].worksheets:
        assert "not real patient or provider performance" in str(ws["A1"].value)
        assert ws.tables, ws.title


def test_rates_are_formulas_and_counts_are_values(workbook) -> None:
    ws = workbook[1]["Executive_Summary"]
    header = {c.value: c.column_letter for c in ws[HEADER_ROW]}
    row = HEADER_ROW + 1
    assert ws[f"{header['Beneficiaries']}{row}"].value == 12
    assert str(ws[f"{header['Admissions per 1000 member years']}{row}"].value).startswith("=IF(")
    assert str(ws[f"{header['Readmission proxy']}{row}"].value).startswith("=IF(")
    assert ws[f"{header['Admissions']}{row}"].value == 8
    util = workbook[1]["Monthly_Utilization"]
    assert any(str(c.value).startswith("=IF(") for c in util[HEADER_ROW + 1])


def test_stored_values_reconcile_to_duckdb(workbook, warehouse) -> None:
    kpis = compute_kpis(warehouse)
    ws = workbook[1]["Executive_Summary"]
    header = {c.value: c.column_letter for c in ws[HEADER_ROW]}
    for i, year in enumerate(sorted(kpis["beneficiaries"])):
        r = HEADER_ROW + 1 + i
        assert ws[f"{header['Year']}{r}"].value == int(year)
        assert ws[f"{header['Claims']}{r}"].value == kpis["claims_total"][year]
        assert ws[f"{header['Paid amount']}{r}"].value == pytest.approx(kpis["payment_total"][year], abs=0.005)
        assert ws[f"{header['Admissions']}{r}"].value == kpis["admissions"][year]
    monthly = workbook[1]["Monthly_Payments"]
    header = {c.value: c.column_letter for c in monthly[HEADER_ROW]}
    total = sum(monthly[f"{header['Payment_amount']}{r}"].value or 0 for r in range(HEADER_ROW + 1, monthly.max_row + 1))
    assert total == pytest.approx(kpis["payment_total_all_years"], abs=0.005)


def test_formula_references_resolve_to_existing_sheets(workbook) -> None:
    wb = workbook[1]
    formulas = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 200
    referenced = {m for f in formulas for m in re.findall(r"([A-Za-z_]+)!", f)}
    assert referenced and referenced <= set(wb.sheetnames)


def test_reconciliation_sheet_ties_workbook_totals_to_duckdb(workbook) -> None:
    ws = workbook[1]["Reconciliation"]
    header = {c.value: c.column_letter for c in ws[HEADER_ROW]}
    assert list(header) == ["Check", "DuckDB", "Tolerance", "Workbook_value", "Variance", "Status"]
    first = HEADER_ROW + 1
    checks = [ws[f"A{r}"].value for r in range(first, first + 10)]
    assert "Monthly paid amount" in checks and "Executive: member years x 12" in checks
    assert all(str(ws[f"{header['Workbook_value']}{r}"].value).startswith("=SUM(") for r in range(first, first + 10))
    assert ws[f"B{first + 11}"].value.startswith("=IF(COUNTIF(F")


def test_no_beneficiary_level_rows_are_exported(workbook, warehouse) -> None:
    ids = {r[0] for r in warehouse.execute("select beneficiary_id from dim_beneficiary").fetchall()}
    for ws in workbook[1].worksheets:
        for row in ws.iter_rows(values_only=True):
            assert not ({str(v) for v in row if v is not None} & ids), ws.title


def test_workbook_has_no_macros(workbook) -> None:
    with zipfile.ZipFile(workbook[0]) as package:
        assert not any("vbaProject" in n for n in package.namelist())


def test_review_list_is_labeled_synthetic(workbook) -> None:
    ws = workbook[1]["Provider_Review"]
    text = " ".join(str(c.value) for row in ws.iter_rows(max_row=3) for c in row if c.value)
    assert "SYNTHETIC" in text and "not fraud" in text
