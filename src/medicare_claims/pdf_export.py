"""Export the Executive_Summary sheet of the Excel workbook to reports/claims_executive_summary.pdf.

LibreOffice renders the PDF headlessly from a one-sheet copy of the workbook, so the PDF uses the same
print settings (landscape, two pages wide, repeated Year column, header and footer) a reader gets in Excel.
The page count and text are checked by tests; the pixels are reviewed by hand (reports/visual_qa.md).
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

SHEET = "Executive_Summary"


def soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def export_executive_pdf(workbook: Path, output: Path) -> Path:
    binary = soffice()
    if binary is None:
        raise RuntimeError("LibreOffice (soffice) is required to export the executive PDF")
    with tempfile.TemporaryDirectory() as tmp:
        single = Path(tmp) / "claims_executive_summary.xlsx"
        wb = load_workbook(workbook)
        for name in list(wb.sheetnames):
            if name != SHEET:
                del wb[name]
        wb.save(single)
        profile = Path(tmp) / "lo-profile"
        subprocess.run([binary, f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
                        "--convert-to", "pdf", "--outdir", tmp, str(single)],
                       check=True, capture_output=True, timeout=180)
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(Path(tmp) / "claims_executive_summary.pdf", output)
    return output
